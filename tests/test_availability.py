"""
Testy pro dostupnost (kousek C):
  - availability_check.detect_soft_404 / detect_bot_challenge / http_status_from_error
  - availability_check.check_not_found_page (mockovaná síť)
  - structure_check: SOFT_404 a EMPTY_HREF
  - links_check: přesměrování (probe_url, check_resources, is_trivial_redirect)
  - main: build_sitemap_report, validate_pages s bot ochranou, run_domain_checks
  - report_excel / report_json: nové sekce a klíče
"""
import io
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from availability_check import (detect_soft_404, detect_soft_404_html, detect_bot_challenge,
                                bot_challenge_message, http_status_from_error,
                                check_not_found_page, SOFT_404_MAX_TEXT_LEN)
from issues import Issue, IssueType
from structure_check import check_structure, _find_empty_links
from stats import page_score
import links_check
from links_check import probe_url, check_resources, is_trivial_redirect

BASE = "https://example.cz/"

ANUBIS_HTML = """<!DOCTYPE html><html><head><title>Making sure you&#39;re not a bot!</title>
<meta name="robots" content="noindex,nofollow"></head><body>
<div id="anubis_challenge" data-x="1"></div><script src="/.within.website/x/cmd/anubis/static/js/main.mjs"></script>
</body></html>"""

CF_HTML = """<html><head><title>Just a moment...</title></head><body>
<div id="cf-challenge-running"></div></body></html>"""


def _page(title="Kontakt | Firma", h1="Kontakt", body="<p>Zavolejte nám.</p>"):
    return (f"<!DOCTYPE html><html lang='cs'><head><title>{title}</title>"
            f"<meta name='viewport' content='width=device-width'>"
            f"<meta name='description' content='Popis stránky s dostatečnou délkou pro test.'>"
            f"</head><body><h1>{h1}</h1>{body}</body></html>")


def _soup(html):
    return BeautifulSoup(html, "html.parser")


def _result(url, category="ok", **kw):
    r = {"url": url, "w3c_category": category, "w3c_warnings": [], "w3c_errors": [],
         "w3c_error_msg": "", "structure_issues": [], "homepage_meta": [], "title": "T",
         "refs": {}, "final_url": ""}
    r.update(kw)
    return r


def _sheet_values(path: Path) -> list[str]:
    ws = load_workbook(path).worksheets[0]
    return [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]


# ── Soft 404 ─────────────────────────────────────────────────────────────────

class TestDetectSoft404(unittest.TestCase):
    POSITIVE = [
        ("Stránka nenalezena | Reality Nadin", "Něco"),
        ("404 – Stránka nebyla nalezena", "Jejda"),
        ("Chyba 404 | Firma", "Text"),
        ("Firma s.r.o.", "404"),
        ("Firma s.r.o.", "Stránka nenalezena"),
        ("Firma", "Požadovaná stránka neexistuje"),
        ("Page not found | Company", "x"),
        ("Company", "Oops! Page not found"),
        ("Firma", "Tato stránka nebyla nalezena."),
        ("Error 404 - Not Found", "x"),
        ("Firma", "Stránku jsme nenašli"),
        ("Firma", "Seite nicht gefunden"),
        ("Firma: Stránka nenalezena", "x"),
        ("Stránka nenalezena (404) | Firma", "Byt 3+kk"),
    ]
    NEGATIVE = [
        ("Prodej bytu 4+kk, 404 m² | Reality", "Prodej bytu 4+kk"),
        ("Prodej pozemku | 404 m²", "Pozemek"),
        ("Jak nastavit vlastní 404 stránku | Blog", "Jak nastavit vlastní 404 stránku"),
        ("Nenalezeny žádné nemovitosti | Reality", "Nabídka nemovitostí"),
        ("Kontakt | Firma", "Kontakt"),
        ("Byt č. 404 – Rezidence Park", "Byt č. 404"),
        ("Firma", "Nabídka již neexistuje"),
        ("O nás – Firma", "Naše historie"),
        ("Reality Nadin – Byty, domy, pozemky", "Vítejte"),
        ("Co dělat, když stránka neexistuje | Nápověda", "Co dělat, když stránka neexistuje"),
        ("Nabídka 4040 | Reality", "Domy 4040"),
    ]

    def test_positive_titles_and_h1(self):
        for title, h1 in self.POSITIVE:
            with self.subTest(title=title, h1=h1):
                self.assertTrue(detect_soft_404(_soup(_page(title, h1))))

    def test_negative_everyday_pages(self):
        for title, h1 in self.NEGATIVE:
            with self.subTest(title=title, h1=h1):
                self.assertEqual(detect_soft_404(_soup(_page(title, h1))), "")

    def test_returns_matching_text(self):
        hit = detect_soft_404(_soup(_page("Firma", "Stránka nenalezena")))
        self.assertEqual(hit, "Stránka nenalezena")

    def test_long_article_about_404_not_flagged(self):
        body = "<p>" + "slovo " * (SOFT_404_MAX_TEXT_LEN // 5) + "</p>"
        self.assertEqual(detect_soft_404(_soup(_page("404 – Stránka nenalezena", "404", body))), "")

    def test_svg_title_in_body_ignored(self):
        html = ("<html><head><title>Kontakt | Firma</title></head><body><h1>Kontakt</h1>"
                "<svg><title>404 ikona</title></svg></body></html>")
        self.assertEqual(detect_soft_404(_soup(html)), "")

    def test_html_variant_and_empty(self):
        self.assertEqual(detect_soft_404_html(""), "")
        self.assertTrue(detect_soft_404_html(_page("Firma", "Page not found")))

    def test_check_structure_reports_soft_404(self):
        issues = check_structure(_page("Stránka nenalezena | Firma", "Stránka nenalezena"),
                                 page_url=BASE)
        soft = [i for i in issues if i.type == IssueType.SOFT_404]
        self.assertEqual(len(soft), 1)
        self.assertIn("Stránka nenalezena", soft[0].items[0])

    def test_check_structure_clean_page_no_soft_404(self):
        issues = check_structure(_page(), page_url=BASE)
        self.assertFalse([i for i in issues if i.type == IssueType.SOFT_404])


# ── Prázdné odkazy ──────────────────────────────────────────────────────────

class TestEmptyLinks(unittest.TestCase):
    def _links(self, inner):
        return _find_empty_links(_soup(f"<html><body>{inner}</body></html>"))

    def test_naked_hash_and_empty_and_javascript(self):
        found = self._links('<a href="#">Obchodní podmínky</a>'
                            '<a href="">Kontakt</a>'
                            '<a href="javascript:void(0)">Více</a>'
                            '<a href="javascript:;">Ještě</a>')
        self.assertEqual(len(found), 4)
        self.assertIn('„Obchodní podmínky“  (href="#")', found)

    def test_anchor_and_real_links_are_fine(self):
        self.assertEqual(self._links('<a href="#kontakt">Kontakt</a>'
                                     '<a href="/o-nas">O nás</a>'
                                     '<a name="top">Nahoru</a>'), [])

    def test_js_hooks_not_reported(self):
        cases = ['<a href="#" class="btn">Menu</a>',
                 '<a href="#" id="open">Menu</a>',
                 '<a href="#" data-bs-toggle="dropdown">Menu</a>',
                 '<a href="#" role="button">Menu</a>',
                 '<a href="#" aria-expanded="false">Menu</a>',
                 '<a href="#" onclick="go()">Menu</a>',
                 '<a href="#" @click="open = true">Menu</a>',
                 '<a href="#" v-on:click="x">Menu</a>',
                 '<a href="#" tabindex="0">Menu</a>']
        for html in cases:
            with self.subTest(html=html):
                self.assertEqual(self._links(html), [])

    def test_image_only_link_not_reported(self):
        self.assertEqual(self._links('<a href="#"><img src="a.png" alt=""></a>'), [])

    def test_submenu_parent_not_reported(self):
        self.assertEqual(self._links('<ul><li><a href="#">Služby</a><ul><li>'
                                     '<a href="/a">A</a></li></ul></li></ul>'), [])

    def test_dedup_and_truncate(self):
        found = self._links('<a href="#">Více</a><a href="#">Více</a>'
                            '<a href="#">' + "x" * 80 + '</a>')
        self.assertEqual(len(found), 2)
        self.assertTrue(any("…" in f for f in found))

    def test_check_structure_issue_and_penalty(self):
        html = _page(body='<p>Text</p><a href="#">Obchodní podmínky</a><a href="#">GDPR</a>')
        issues = check_structure(html, page_url=BASE)
        empty = [i for i in issues if i.type == IssueType.EMPTY_HREF]
        self.assertEqual(len(empty), 1)
        self.assertEqual(empty[0].total_count, 2)
        clean = page_score({"w3c_category": "ok", "w3c_errors": [],
                            "structure_issues": []})
        dirty = page_score({"w3c_category": "ok", "w3c_errors": [],
                            "structure_issues": [empty[0],
                                                 Issue(type=IssueType.SOFT_404, items=["x"])]})
        self.assertEqual(clean - dirty, 2 + 10)


# ── Bot ochrana ─────────────────────────────────────────────────────────────

class TestBotChallenge(unittest.TestCase):
    def test_anubis_detected(self):
        self.assertEqual(detect_bot_challenge(ANUBIS_HTML), "Anubis")
        self.assertEqual(detect_bot_challenge(
            "<html><title>x</title><script>techaro.lol-anubis-cookie</script></html>"), "Anubis")

    def test_cloudflare_detected(self):
        self.assertEqual(detect_bot_challenge(CF_HTML), "Cloudflare")
        self.assertEqual(detect_bot_challenge(
            "<html><head><title>Attention Required! | Cloudflare</title></head></html>"),
            "Cloudflare")

    def test_plain_page_not_flagged(self):
        html = ("<html><head><title>Just a moment to relax | Wellness</title></head>"
                "<body><p>We are not a bot friendly spa. Cloudflare hosts us.</p></body></html>")
        self.assertEqual(detect_bot_challenge(html), "")
        self.assertEqual(detect_bot_challenge(""), "")
        self.assertEqual(detect_bot_challenge(_page()), "")

    def test_message_mentions_provider(self):
        self.assertIn("Anubis", bot_challenge_message("Anubis"))


class TestHttpStatusFromError(unittest.TestCase):
    def test_parses_requests_messages(self):
        self.assertEqual(http_status_from_error(
            "404 Client Error: Not Found for url: https://e.cz/x"), 404)
        self.assertEqual(http_status_from_error("503 Server Error: Service Unavailable"), 503)

    def test_network_errors_are_zero(self):
        self.assertEqual(http_status_from_error("HTTPSConnectionPool: Max retries exceeded"), 0)
        self.assertEqual(http_status_from_error(""), 0)
        self.assertEqual(http_status_from_error(None), 0)


# ── Test vlastní 404 stránky ────────────────────────────────────────────────

def _resp(status=200, text="", url="", location=None):
    r = MagicMock()
    r.status_code = status
    r.text = text
    r.url = url
    r.headers = {"Location": location} if location else {}
    return r


class TestCheckNotFoundPage(unittest.TestCase):
    def _run(self, responses):
        """responses = seznam odpovědí v pořadí volání session.get."""
        s = MagicMock()
        s.get.side_effect = responses
        return check_not_found_page(BASE, session=s), s

    def test_404_is_ok(self):
        res, s = self._run([_resp(404)])
        self.assertEqual(res["verdict"], "ok")
        self.assertEqual(res["status"], 404)
        self.assertIn("HTTP 404", res["message"])
        self.assertFalse(s.get.call_args.kwargs.get("allow_redirects"))
        self.assertTrue(res["url"].startswith("https://example.cz/wv-neexistujici-stranka-"))

    def test_410_is_ok(self):
        res, _ = self._run([_resp(410)])
        self.assertEqual(res["verdict"], "ok")

    def test_200_plain_is_soft_404(self):
        res, _ = self._run([_resp(200, text=_page("Firma", "Vítejte"))])
        self.assertEqual(res["verdict"], "soft_404")
        self.assertEqual(res["soft_404_text"], "")
        self.assertIn("HTTP 200", res["message"])

    def test_200_with_404_text_reports_wrong_status(self):
        res, _ = self._run([_resp(200, text=_page("Firma", "Stránka nenalezena"))])
        self.assertEqual(res["verdict"], "soft_404")
        self.assertEqual(res["soft_404_text"], "Stránka nenalezena")
        self.assertIn("soft 404", res["message"])

    def test_redirect_to_homepage(self):
        res, s = self._run([_resp(301, location="/"),
                            _resp(200, url="https://example.cz/", text=_page())])
        self.assertEqual(res["verdict"], "redirect_home")
        self.assertEqual(res["final_url"], "https://example.cz/")
        self.assertEqual(s.get.call_count, 2)

    def test_redirect_to_404_page_is_ok(self):
        res, _ = self._run([_resp(302, location="/404"),
                            _resp(404, url="https://example.cz/404")])
        self.assertEqual(res["verdict"], "ok")
        self.assertIn("/404", res["message"])

    def test_redirect_to_other_200_page(self):
        res, _ = self._run([_resp(302, location="/chyba/"),
                            _resp(200, url="https://example.cz/chyba/",
                                  text=_page("Firma", "Stránka nenalezena"))])
        self.assertEqual(res["verdict"], "redirect_200")
        self.assertIn("Stránka nenalezena", res["message"])

    def test_server_error_forbidden_unknown(self):
        self.assertEqual(self._run([_resp(500)])[0]["verdict"], "server_error")
        self.assertEqual(self._run([_resp(403)])[0]["verdict"], "forbidden")
        self.assertEqual(self._run([_resp(204)])[0]["verdict"], "unknown")

    def test_connection_error_unreachable(self):
        res, _ = self._run([ConnectionError("boom")])
        self.assertEqual(res["verdict"], "unreachable")
        self.assertEqual(res["status"], 0)
        self.assertIn("ConnectionError", res["message"])

    def test_redirect_target_unreachable(self):
        res, _ = self._run([_resp(301, location="/x"), ConnectionError("boom")])
        self.assertEqual(res["verdict"], "unknown")


# ── Přesměrování v links_check ───────────────────────────────────────────────

class TestRedirects(unittest.TestCase):
    def test_probe_url_records_redirect(self):
        s = MagicMock()
        first = _resp(301, location="/nova/")
        final = _resp(200, url="https://example.cz/nova/")
        final.headers = {"Content-Length": "10"}
        s.head.side_effect = [first, final]
        info = probe_url(s, "https://example.cz/stara/")
        self.assertEqual(info["status"], 200)
        self.assertEqual(info["redirect"], "https://example.cz/nova/")
        self.assertEqual(info["redirect_status"], 301)
        self.assertEqual(s.head.call_count, 2)
        self.assertTrue(s.head.call_args.kwargs.get("allow_redirects"))

    def test_probe_url_redirect_location_fallback(self):
        s = MagicMock()
        first = _resp(302, location="../cil")
        final = _resp(200)            # mock bez str url → použije se Location
        s.head.side_effect = [first, final]
        info = probe_url(s, "https://example.cz/a/b/")
        self.assertEqual(info["redirect"], "https://example.cz/a/cil")

    def test_probe_url_redirect_then_404_is_broken(self):
        s = MagicMock()
        s.head.side_effect = [_resp(301, location="/x"), _resp(404, url="https://example.cz/x")]
        info = probe_url(s, "https://example.cz/a")
        self.assertEqual(info["status"], 404)
        self.assertEqual(info["redirect"], "https://example.cz/x")

    def test_is_trivial_redirect(self):
        self.assertTrue(is_trivial_redirect("http://example.cz/a", "https://www.example.cz/a/"))
        self.assertFalse(is_trivial_redirect("https://example.cz/a", "https://example.cz/b"))

    def test_check_resources_collects_internal_redirects(self):
        table = {
            "https://example.cz/stara/": {"status": 200, "size": None, "error": "",
                                          "redirect": "https://example.cz/nova/", "redirect_status": 301},
            "http://example.cz/kontakt": {"status": 200, "size": None, "error": "",
                                          "redirect": "https://example.cz/kontakt/", "redirect_status": 301},
            "https://jiny.cz/x": {"status": 200, "size": None, "error": "",
                                  "redirect": "https://jiny.cz/y", "redirect_status": 302},
            "https://example.cz/ok": {"status": 200, "size": None, "error": ""},
            # přesměrování na 404 → je v broken_links, v redirects už ne
            "https://example.cz/mrtvy": {"status": 404, "size": None, "error": "",
                                         "redirect": "https://example.cz/mrtvy/", "redirect_status": 301},
        }
        results = [_result(BASE, refs={"links": list(table), "images": []})]
        with patch.object(links_check, "probe_url", side_effect=lambda s, u, timeout=0: table[u]), \
             patch.object(links_check, "make_session", return_value=MagicMock()):
            report = check_resources(results, BASE, check_external=True)
        redirects = report["redirects"]
        self.assertEqual([r["url"] for r in redirects],
                         ["https://example.cz/stara/", "http://example.cz/kontakt"])
        self.assertFalse(redirects[0]["trivial"])
        self.assertTrue(redirects[1]["trivial"])
        self.assertEqual(redirects[0]["to"], "https://example.cz/nova/")
        self.assertEqual(redirects[0]["sources"], [BASE])
        self.assertEqual([b["url"] for b in report["broken_links"]], ["https://example.cz/mrtvy"])
        # přesměrování není Issue → jediné Issue je nefunkční odkaz
        self.assertEqual([i.type for i in results[0]["structure_issues"]], [IssueType.BROKEN_LINK])


# ── main.py: sitemap report, bot ochrana ve validaci, doménové kontroly ─────

class TestBuildSitemapReport(unittest.TestCase):
    def test_none_without_sitemap(self):
        import main as m
        self.assertIsNone(m.build_sitemap_report([], [_result(BASE)]))

    def test_dead_redirected_unreachable(self):
        import main as m
        sm = [BASE, BASE + "a", BASE + "b", BASE + "c", BASE + "d"]
        results = [
            _result(BASE),
            _result(BASE + "a", "validator_error",
                    w3c_error_msg="404 Client Error: Not Found for url: x", http_status=404),
            _result(BASE + "b", "validator_error",
                    w3c_error_msg="HTTPSConnectionPool: Max retries exceeded", http_status=0),
            _result(BASE + "c", final_url=BASE + "c-nova/"),
            _result(BASE + "d", "validator_error", bot_challenge="Anubis",
                    w3c_error_msg="Blokováno…", http_status=200),
            _result(BASE + "mimo-sitemapu", "validator_error",
                    w3c_error_msg="404 Client Error: Not Found", http_status=404),
        ]
        rep = m.build_sitemap_report(sm, results)
        self.assertEqual(rep["total"], 5)
        self.assertEqual([e["url"] for e in rep["dead"]], [BASE + "a"])
        self.assertEqual(rep["dead"][0]["status"], 404)
        self.assertEqual([e["url"] for e in rep["unreachable"]], [BASE + "b"])
        self.assertEqual(rep["redirected"], [{"url": BASE + "c", "to": BASE + "c-nova/",
                                              "trivial": False}])

    def test_status_parsed_from_message_when_missing(self):
        import main as m
        results = [_result(BASE + "a", "validator_error",
                           w3c_error_msg="410 Client Error: Gone for url: x")]
        rep = m.build_sitemap_report([BASE + "a"], results)
        self.assertEqual(rep["dead"][0]["status"], 410)


class TestValidatePagesBotChallenge(unittest.TestCase):
    def test_challenge_page_not_validated(self):
        import main as m
        fetched = {BASE: (ANUBIS_HTML.encode(), ANUBIS_HTML, "text/html", ""),
                   BASE + "ok": (_page().encode(), _page(), "text/html", "")}
        validate = MagicMock(return_value={"category": "ok", "warnings": [], "errors": [],
                                           "error_msg": ""})
        with patch.object(m, "fetch_html", side_effect=lambda s, u: fetched[u]), \
             patch.object(m.w3c_mod, "validate", validate), \
             redirect_stdout(io.StringIO()) as out:
            results = m.validate_pages([BASE, BASE + "ok"], start_url=BASE, delay=0)
        blocked = results[0]
        self.assertEqual(blocked["w3c_category"], "validator_error")
        self.assertEqual(blocked["bot_challenge"], "Anubis")
        self.assertIn("Anubis", blocked["w3c_error_msg"])
        self.assertEqual(results[1]["w3c_category"], "ok")
        # W3C validátor se pro challenge stránku nespustil
        self.assertEqual(validate.call_count, 1)
        self.assertIn("BOT OCHRANA", out.getvalue())
        self.assertIn("NENÍ platný", out.getvalue())

    def test_failed_fetch_keeps_http_status(self):
        import main as m
        with patch.object(m, "fetch_html",
                          return_value=(None, None, "404 Client Error: Not Found for url: x", "")), \
             redirect_stdout(io.StringIO()):
            results = m.validate_pages([BASE + "x"], start_url=BASE, delay=0)
        self.assertEqual(results[0]["http_status"], 404)
        self.assertEqual(results[0]["w3c_category"], "validator_error")


class TestRunDomainChecksNotFound(unittest.TestCase):
    def test_not_found_included(self):
        import main as m
        nf = {"url": BASE + "wv-x/", "status": 404, "final_status": 404, "final_url": "",
              "verdict": "ok", "message": "ok", "soft_404_text": ""}
        with patch.object(m, "check_robots_js_css", return_value=([], False)), \
             patch.object(m, "check_user_pages", return_value=[]), \
             patch.object(m, "check_not_found_page", return_value=nf):
            info = m.run_domain_checks(BASE)
        self.assertEqual(info["not_found"], nf)

    def test_not_found_failure_is_contained(self):
        import main as m
        with patch.object(m, "check_robots_js_css", return_value=([], False)), \
             patch.object(m, "check_user_pages", return_value=[]), \
             patch.object(m, "check_not_found_page", side_effect=RuntimeError("x")):
            info = m.run_domain_checks(BASE)
        self.assertEqual(info["not_found"]["verdict"], "unreachable")


# ── Excel + JSON ─────────────────────────────────────────────────────────────

class TestReportsAvailability(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name) / "r.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def _write(self, results, **kw):
        from report_excel import write_report
        write_report(results, self.out, BASE, score=80, **kw)
        return _sheet_values(self.out)

    def test_bot_blocked_warning_first_in_summary(self):
        results = [_result(BASE, "validator_error", bot_challenge="Anubis",
                           w3c_error_msg=bot_challenge_message("Anubis"))]
        vals = self._write(results)
        joined = "\n".join(vals)
        self.assertIn("AUDIT NENÍ PLATNÝ", joined)
        self.assertIn("Anubis", joined)
        self.assertIn("Zablokováno bot ochranou (stránek)", vals)
        # varování je hned pod nadpisem SOUHRN
        self.assertLess(vals.index("SOUHRN"),
                        next(i for i, v in enumerate(vals) if "AUDIT NENÍ PLATNÝ" in v))
        self.assertLess(next(i for i, v in enumerate(vals) if "AUDIT NENÍ PLATNÝ" in v),
                        vals.index("Web Quality Score"))

    def test_no_bot_warning_normally(self):
        vals = self._write([_result(BASE)])
        self.assertFalse(any("AUDIT NENÍ PLATNÝ" in v for v in vals))

    def test_sitemap_section(self):
        rep = {"total": 3, "dead": [{"url": BASE + "a", "status": 404, "error": "404"}],
               "unreachable": [], "redirected": [{"url": BASE + "c", "to": BASE + "c2"}]}
        vals = self._write([_result(BASE)], sitemap_report=rep)
        self.assertIn("SITEMAP.XML – NEEXISTUJÍCÍ / PŘESMĚROVANÉ URL", vals)
        self.assertIn("HTTP 404", vals)
        self.assertIn(f"→ {BASE}c2", vals)
        self.assertTrue(any("1 neexistujících" in v and "1 přesměrovaných" in v for v in vals))

    def test_sitemap_section_ok_and_absent(self):
        vals = self._write([_result(BASE)],
                           sitemap_report={"total": 4, "dead": [], "unreachable": [], "redirected": []})
        self.assertTrue(any("Všech 4 URL ze sitemapy" in v for v in vals))
        vals = self._write([_result(BASE)], sitemap_report=None)
        self.assertNotIn("SITEMAP.XML – NEEXISTUJÍCÍ / PŘESMĚROVANÉ URL", vals)

    def test_not_found_section(self):
        nf = {"url": BASE + "wv-x/", "status": 200, "final_status": 200, "final_url": "",
              "verdict": "soft_404", "soft_404_text": "",
              "message": "Web vrací HTTP 200 pro neexistující stránku (soft 404)"}
        vals = self._write([_result(BASE)], domain_info={"not_found": nf})
        self.assertIn("NEEXISTUJÍCÍ STRÁNKA – TEST HTTP 404", vals)
        self.assertTrue(any(v.startswith("⚠ Web vrací HTTP 200") for v in vals))
        self.assertTrue(any("testovaná URL" in v and "HTTP 200" in v for v in vals))
        vals = self._write([_result(BASE)], domain_info={"not_found": dict(nf, verdict="ok",
                                                                            message="OK 404")})
        self.assertIn("✓ OK 404", vals)
        vals = self._write([_result(BASE)], domain_info={})
        self.assertIn("Kontrola neproběhla nebo nedostupná", vals)

    def test_redirects_in_links_section(self):
        link_report = {"broken_links": [], "images": [], "checked_links": 3, "checked_images": 0,
                       "known_ok": 0, "skipped_external": 0, "collapsed_query": 0,
                       "skipped_limit": 0, "skipped_time": 0, "unverified": 0, "aborted": "",
                       "elapsed": 0.1, "check_external": False,
                       "redirects": [
                           {"url": BASE + "stara/", "status": 301, "to": BASE + "nova/",
                            "final_status": 200, "trivial": False, "sources": [BASE]},
                           {"url": "http://example.cz/k", "status": 301, "to": BASE + "k/",
                            "final_status": 200, "trivial": True, "sources": [BASE]},
                       ]}
        vals = self._write([_result(BASE)], link_report=link_report)
        self.assertTrue(any("Odkazy vedoucí přes přesměrování: 2" in v
                            and "z toho 1 jen http→https" in v for v in vals))
        self.assertIn(BASE + "nova/", vals)
        self.assertIn("301", vals)

    def test_json_keys(self):
        from report_json import build_json
        nf = {"verdict": "ok", "message": "m"}
        results = [_result(BASE, "validator_error", bot_challenge="Anubis", http_status=200,
                           final_url=BASE + "x")]
        data = build_json(results, BASE, domain_info={"not_found": nf},
                          link_report={"redirects": [{"url": "a"}]},
                          sitemap_report={"total": 1, "dead": [], "unreachable": [], "redirected": []})
        self.assertEqual(data["domain"]["not_found"], nf)
        self.assertEqual(data["redirects"], [{"url": "a"}])
        self.assertEqual(data["sitemap"]["total"], 1)
        page = data["pages"][0]
        self.assertEqual(page["bot_challenge"], "Anubis")
        self.assertEqual(page["http_status"], 200)
        self.assertEqual(page["final_url"], BASE + "x")

    def test_json_defaults(self):
        from report_json import build_json
        data = build_json([_result(BASE)], BASE)
        self.assertIsNone(data["sitemap"])
        self.assertEqual(data["redirects"], [])
        self.assertIsNone(data["domain"]["not_found"])
        self.assertEqual(data["pages"][0]["bot_challenge"], "")


if __name__ == "__main__":
    unittest.main()
