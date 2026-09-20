"""
Testy pro SEO modul za přepínačem `--seo` (kousek D):
  - issues.SEO_ISSUE_TYPES / is_seo_issue – konzistence s vahami a popisky
  - structure_check.check_structure(seo=) – výchozí stav SEO nálezy nevrací
  - stats – skóre bez SEO se počítá jen z jádra
  - main.validate_pages(seo=) – homepage meta a duplicitní <title> jen se seo
  - main.run_link_checks(seo=) – předání do check_resources
  - report_excel – sekce SEO / META – HOMEPAGE jen se seo, poznámka bez něj
  - report_json – klíč "seo", compare_runs při rozdílném režimu
"""
import io
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

from issues import Issue, IssueType, ISSUE_LABELS, SEO_ISSUE_TYPES, is_seo_issue
from structure_check import check_structure
from stats import page_score, compute_stats, _BINARY_PENALTIES, _COUNTED_PENALTIES
from report_json import build_json, compare_runs

BASE = "https://example.cz/"

# Stránka, která má SEO problémy všeho druhu, ale jádro v pořádku
# (jen H1, title, lang, viewport – žádné prázdné tagy ani testovací obsah).
SEO_BAD_HTML = """<!DOCTYPE html><html lang="cs"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width"><title>Firma – úvodní stránka webu</title>
</head><body>
<h1>Nadpis</h1><h3>Přeskočený nadpis</h3>
<img src="/a.jpg">
<a href="https://jinde.cz/">Externí bez noopener</a>
<p>Běžný text stránky o firmě a jejích službách.</p>
</body></html>"""

# Stránka s problémy jádra i SEO
MIXED_HTML = """<!DOCTYPE html><html><head><meta charset="utf-8"><title>T</title></head>
<body><h1>A</h1><h1>B</h1><div></div><img src="/a.jpg"></body></html>"""


def _sheet_values(path: Path) -> list[str]:
    wb = load_workbook(path)
    vals = []
    for row in wb.worksheets[0].iter_rows(values_only=True):
        for v in row:
            if v is not None:
                vals.append(str(v))
    return vals


def _page(url=BASE, issues=None, homepage_meta=None, title="T"):
    return {"url": url, "w3c_category": "ok", "w3c_warnings": [], "w3c_errors": [],
            "w3c_error_msg": "", "structure_issues": issues or [],
            "homepage_meta": homepage_meta or [], "title": title}


class TestSeoIssueTypes(unittest.TestCase):
    def test_expected_members(self):
        expected = {IssueType.HEADING_SKIP, IssueType.MISSING_META_DESC,
                    IssueType.EMPTY_META_DESC, IssueType.MISSING_ALT,
                    IssueType.EXTERNAL_LINK, IssueType.DUPLICATE_TITLE,
                    IssueType.MISSING_CANONICAL, IssueType.CANONICAL_MISMATCH,
                    IssueType.CANONICAL_HTTP, IssueType.MISSING_OG,
                    IssueType.IMG_NO_DIMENSIONS, IssueType.IMG_TOO_LARGE}
        self.assertEqual(set(SEO_ISSUE_TYPES), expected)

    def test_core_types_are_not_seo(self):
        for t in (IssueType.NOINDEX, IssueType.STAGING_URL, IssueType.MISSING_H1,
                  IssueType.MISSING_TITLE, IssueType.MISSING_LANG, IssueType.MISSING_VIEWPORT,
                  IssueType.FORBIDDEN_CONTENT, IssueType.TEMPLATE_VARIABLE,
                  IssueType.DEFAULT_META_TEXT, IssueType.SOFT_404, IssueType.EMPTY_HREF,
                  IssueType.BROKEN_LINK, IssueType.IMG_BROKEN, IssueType.HTTP_LINK,
                  IssueType.EMPTY_TAG, IssueType.DUPLICATE_ID):
            self.assertNotIn(t, SEO_ISSUE_TYPES, t)

    def test_every_seo_type_has_penalty_and_label(self):
        for t in SEO_ISSUE_TYPES:
            self.assertTrue(t in _BINARY_PENALTIES or t in _COUNTED_PENALTIES, t)
            self.assertIn(t, ISSUE_LABELS)

    def test_is_seo_issue_accepts_issue_and_type(self):
        self.assertTrue(is_seo_issue(Issue(type=IssueType.MISSING_OG)))
        self.assertTrue(is_seo_issue(IssueType.MISSING_ALT))
        self.assertFalse(is_seo_issue(Issue(type=IssueType.MISSING_H1)))
        self.assertFalse(is_seo_issue("nesmysl"))


class TestCheckStructureSwitch(unittest.TestCase):
    def test_default_hides_seo_issues(self):
        issues = check_structure(SEO_BAD_HTML, page_url=BASE)
        self.assertEqual(issues, [], [i.type for i in issues])

    def test_seo_true_reports_them(self):
        types = {i.type for i in check_structure(SEO_BAD_HTML, page_url=BASE, seo=True)}
        for t in (IssueType.HEADING_SKIP, IssueType.MISSING_META_DESC, IssueType.MISSING_ALT,
                  IssueType.EXTERNAL_LINK, IssueType.MISSING_CANONICAL, IssueType.MISSING_OG,
                  IssueType.IMG_NO_DIMENSIONS):
            self.assertIn(t, types)
        self.assertTrue(types <= SEO_ISSUE_TYPES, types - SEO_ISSUE_TYPES)

    def test_core_issues_survive_in_both_modes(self):
        core = {IssueType.MULTIPLE_H1, IssueType.EMPTY_TAG, IssueType.MISSING_LANG,
                IssueType.MISSING_VIEWPORT}
        off = {i.type for i in check_structure(MIXED_HTML, page_url=BASE)}
        on  = {i.type for i in check_structure(MIXED_HTML, page_url=BASE, seo=True)}
        self.assertTrue(core <= off, core - off)
        self.assertTrue(core <= on, core - on)
        self.assertFalse(off & SEO_ISSUE_TYPES)
        self.assertTrue(on & SEO_ISSUE_TYPES)
        # bez SEO = podmnožina se SEO
        self.assertTrue(off <= on)

    def test_score_without_seo_is_core_only(self):
        r_off = _page(issues=check_structure(SEO_BAD_HTML, page_url=BASE))
        r_on  = _page(issues=check_structure(SEO_BAD_HTML, page_url=BASE, seo=True))
        self.assertEqual(page_score(r_off), 100.0)
        self.assertLess(page_score(r_on), 100.0)
        self.assertEqual(compute_stats([r_off]).struct_ok, 1)
        self.assertEqual(compute_stats([r_on]).struct_bad, 1)


class TestValidatePagesSwitch(unittest.TestCase):
    HTML = ("<!DOCTYPE html><html lang='cs'><head><meta charset='utf-8'><title>Stejný titulek"
            "</title><meta name='viewport' content='width=device-width'></head>"
            "<body><h1>X</h1><p>Text</p></body></html>")

    def _run(self, seo):
        import main as m
        pages = [BASE, BASE + "a/"]
        fetched = {u: (self.HTML.encode(), self.HTML, "text/html", "") for u in pages}
        validate = MagicMock(return_value={"category": "ok", "warnings": [], "errors": [],
                                           "error_msg": ""})
        with patch.object(m, "fetch_html", side_effect=lambda s, u: fetched[u]), \
             patch.object(m.w3c_mod, "validate", validate), \
             redirect_stdout(io.StringIO()) as out:
            results = m.validate_pages(pages, start_url=BASE, delay=0, seo=seo)
        return results, out.getvalue()

    def test_seo_off(self):
        results, out = self._run(seo=False)
        types = {i.type for r in results for i in r["structure_issues"]}
        self.assertNotIn(IssueType.DUPLICATE_TITLE, types)
        self.assertNotIn(IssueType.MISSING_META_DESC, types)
        self.assertEqual(results[0]["homepage_meta"], [])
        self.assertNotIn("Duplicitní <title>", out)

    def test_seo_on(self):
        results, out = self._run(seo=True)
        types = {i.type for r in results for i in r["structure_issues"]}
        self.assertIn(IssueType.DUPLICATE_TITLE, types)
        self.assertIn(IssueType.MISSING_META_DESC, types)
        self.assertTrue(results[0]["homepage_meta"])   # homepage = start_url
        self.assertEqual(results[1]["homepage_meta"], [])
        self.assertIn("Duplicitní <title>", out)

    def test_run_link_checks_passes_seo(self):
        import main as m
        report = {"broken_links": [], "images": [], "checked_links": 0, "checked_images": 0,
                  "known_ok": 0, "skipped_external": 0, "elapsed": 0.0}
        with patch.object(m, "check_resources", return_value=report) as cr, \
             redirect_stdout(io.StringIO()) as out:
            m.run_link_checks([_page()], BASE, seo=True)
            self.assertTrue(cr.call_args.kwargs["seo"])
            m.run_link_checks([_page()], BASE)
            self.assertFalse(cr.call_args.kwargs["seo"])
        self.assertIn("velikost jen s --seo", out.getvalue())


class TestReportsSeo(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name) / "r.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def _write(self, results, **kw):
        from report_excel import write_report
        write_report(results, self.out, BASE, score=90, **kw)
        return _sheet_values(self.out)

    def _results(self):
        return [_page(issues=[Issue(type=IssueType.MISSING_H1),
                              Issue(type=IssueType.MISSING_OG, items=["chybí og:title"]),
                              Issue(type=IssueType.MISSING_ALT, items=["/a.jpg"])],
                      homepage_meta=["Title: 45 znaků – v pořádku"])]

    def test_excel_without_seo(self):
        vals = self._write(self._results())
        self.assertTrue(any("SEO kontroly" in v and "--seo" in v for v in vals))
        self.assertFalse(any(v.startswith("SEO – SOUHRN") for v in vals))
        self.assertFalse(any("META – HOMEPAGE" in v for v in vals))
        self.assertIn(ISSUE_LABELS[IssueType.MISSING_H1], vals)
        # SEO nálezy se v HTML struktuře nezobrazí, i kdyby ve výsledcích byly
        self.assertNotIn(ISSUE_LABELS[IssueType.MISSING_OG], vals)
        self.assertNotIn(ISSUE_LABELS[IssueType.MISSING_ALT], vals)

    def test_excel_with_seo(self):
        vals = self._write(self._results(), seo=True)
        self.assertFalse(any("SEO kontroly" in v and "vypnuté" in v for v in vals))
        self.assertTrue(any(v.startswith("SEO – SOUHRN PROBLÉMŮ") for v in vals))
        self.assertTrue(any("META – HOMEPAGE" in v for v in vals))
        self.assertIn(ISSUE_LABELS[IssueType.MISSING_H1], vals)
        self.assertIn(ISSUE_LABELS[IssueType.MISSING_OG], vals)
        self.assertIn(ISSUE_LABELS[IssueType.MISSING_ALT], vals)
        # pořadí: struktura (jádro) před SEO sekcí
        i_struct = vals.index("HTML STRUKTURA – SOUHRN PROBLÉMŮ")
        i_seo = next(i for i, v in enumerate(vals) if v.startswith("SEO – SOUHRN"))
        self.assertLess(i_struct, i_seo)
        self.assertLess(i_seo, vals.index(ISSUE_LABELS[IssueType.MISSING_OG]))

    def test_excel_seo_empty_message(self):
        vals = self._write([_page(issues=[Issue(type=IssueType.MISSING_H1)])], seo=True)
        self.assertIn("✓ Žádné SEO problémy nalezeny", vals)

    def test_json_seo_flag(self):
        self.assertFalse(build_json([_page()], BASE)["seo"])
        self.assertTrue(build_json([_page()], BASE, seo=True)["seo"])

    def test_compare_runs_ignores_seo_when_modes_differ(self):
        prev = build_json([_page(issues=[Issue(type=IssueType.MISSING_OG),
                                         Issue(type=IssueType.MISSING_H1)])], BASE, seo=True)
        cur = build_json([_page(issues=[])], BASE, seo=False)
        cmp_ = compare_runs(prev, cur)
        self.assertTrue(cmp_["seo_ignored"])
        self.assertEqual([f["label"] for f in cmp_["fixed"]],
                         [ISSUE_LABELS[IssueType.MISSING_H1]])
        # opačný směr: SEO problémy „nepřibyly“
        cmp2 = compare_runs(cur, prev)
        self.assertTrue(cmp2["seo_ignored"])
        self.assertEqual([n["label"] for n in cmp2["new"]],
                         [ISSUE_LABELS[IssueType.MISSING_H1]])

    def test_compare_runs_same_mode_keeps_seo(self):
        prev = build_json([_page(issues=[Issue(type=IssueType.MISSING_OG)])], BASE, seo=True)
        cur = build_json([_page(issues=[])], BASE, seo=True)
        cmp_ = compare_runs(prev, cur)
        self.assertFalse(cmp_["seo_ignored"])
        self.assertEqual(cmp_["fixed_count"], 1)

    def test_compare_runs_old_json_without_seo_key_counts_as_seo(self):
        prev = build_json([_page(issues=[Issue(type=IssueType.MISSING_OG)])], BASE, seo=True)
        del prev["seo"]
        cur = build_json([_page(issues=[])], BASE, seo=False)
        self.assertTrue(compare_runs(prev, cur)["seo_ignored"])
        cur_seo = build_json([_page(issues=[])], BASE, seo=True)
        self.assertFalse(compare_runs(prev, cur_seo)["seo_ignored"])

    def test_excel_comparison_note(self):
        prev = build_json([_page(issues=[Issue(type=IssueType.MISSING_OG)])], BASE, seo=True)
        cur = build_json([_page(issues=[])], BASE, seo=False)
        vals = self._write([_page()], comparison=compare_runs(prev, cur))
        self.assertTrue(any("SEO problémy se v porovnání ignorují" in v for v in vals))


if __name__ == "__main__":
    unittest.main()
