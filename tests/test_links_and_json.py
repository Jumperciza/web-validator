"""
Testy pro links_check.py (odkazy a obrázky), report_json.py (JSON export,
porovnání s minulým během) a nové sekce Excel reportu.
Vše bez sítě – HEAD/GET requesty jsou mockované.
"""
import json
import os
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

import links_check
from issues import Issue, IssueType
from links_check import (extract_refs, probe_url, check_resources, url_key,
                         collapse_query, is_network_error)
from report_json import (build_json, write_json, build_json_path,
                         find_previous_json, load_previous, compare_runs,
                         format_previous_date)


BASE = "https://example.cz/"


def _page(url, category="ok", issues=None, refs=None, errors=None, title="T"):
    return {
        "url":              url,
        "w3c_category":     category,
        "w3c_warnings":     [],
        "w3c_errors":       errors or [],
        "w3c_error_msg":    "",
        "structure_issues": issues or [],
        "homepage_meta":    [],
        "title":            title,
        "refs":             refs or {},
    }


def _sheet_values(path: Path) -> list[str]:
    ws = load_workbook(path).worksheets[0]
    return [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]


# ── extract_refs ─────────────────────────────────────────────────────────────

class TestExtractRefs(unittest.TestCase):
    def test_links_absolute_deduped_without_fragment(self):
        html = ('<a href="/o-nas">1</a><a href="/o-nas#tym">2</a>'
                '<a href="https://example.cz/o-nas/">3</a>'
                '<a href="https://jinde.cz/x">ext</a>')
        refs = extract_refs(html, BASE)
        self.assertEqual(refs["links"], ["https://example.cz/o-nas",
                                         "https://example.cz/o-nas/",
                                         "https://jinde.cz/x"])

    def test_non_http_targets_skipped(self):
        html = ('<a href="mailto:a@b.cz">m</a><a href="tel:+420">t</a>'
                '<a href="javascript:void(0)">j</a><a href="#top">a</a>'
                '<a href="">e</a><a href="ftp://x.cz/f">f</a>')
        self.assertEqual(extract_refs(html, BASE)["links"], [])

    def test_images_src_data_src_and_og_image(self):
        html = ('<meta property="og:image" content="/img/og.png">'
                '<img src="/a.png"><img data-src="/lazy.png">'
                '<img src="data:image/png;base64,xxx"><img src="/a.png">')
        refs = extract_refs(html, BASE)
        self.assertEqual(refs["images"], ["https://example.cz/a.png",
                                          "https://example.cz/lazy.png",
                                          "https://example.cz/img/og.png"])

    def test_url_key_ignores_scheme_www_slash(self):
        self.assertEqual(url_key("http://www.example.cz/a/"), url_key("https://example.cz/a"))
        self.assertNotEqual(url_key("https://example.cz/a?p=1"), url_key("https://example.cz/a"))


# ── probe_url ────────────────────────────────────────────────────────────────

def _resp(status=200, headers=None):
    r = MagicMock()
    r.status_code = status
    r.headers = headers or {}
    return r


class TestProbeUrl(unittest.TestCase):
    def test_head_ok_with_size(self):
        s = MagicMock()
        s.head.return_value = _resp(200, {"Content-Length": "2048"})
        info = probe_url(s, "https://example.cz/a.png")
        self.assertEqual(info, {"status": 200, "size": 2048, "error": ""})
        s.get.assert_not_called()

    def test_head_405_falls_back_to_get(self):
        s = MagicMock()
        s.head.return_value = _resp(405)
        s.get.return_value  = _resp(200, {"Content-Length": "abc"})
        info = probe_url(s, "https://example.cz/a")
        self.assertEqual(info["status"], 200)
        self.assertIsNone(info["size"])
        self.assertTrue(s.get.call_args.kwargs.get("stream"))

    def test_exception_is_status_zero(self):
        s = MagicMock()
        s.head.side_effect = ConnectionError("boom")
        info = probe_url(s, "https://example.cz/a")
        self.assertEqual(info["status"], 0)
        self.assertIn("boom", info["error"])

    def test_head_keeps_connection_alive_get_stream_is_closed(self):
        """resp.close() po HEAD by spojení vyhodil z poolu (nové DNS+TLS na
        každý request) – HEAD se jen 'dočte', zavírá se jen GET stream."""
        s = MagicMock()
        head = _resp(200)
        s.head.return_value = head
        probe_url(s, "https://example.cz/a")
        head.close.assert_not_called()

        s.head.return_value = _resp(405)
        get = _resp(200)
        s.get.return_value = get
        probe_url(s, "https://example.cz/b")
        get.close.assert_called_once()


class TestCollapseAndNetworkError(unittest.TestCase):
    def test_query_dropped_on_pretty_paths(self):
        self.assertEqual(collapse_query("https://e.cz/produkty/kremy/?p13%5B0%5D=5&r=cs"),
                         "https://e.cz/produkty/kremy/")
        self.assertEqual(collapse_query("https://e.cz/kontakt?r=cs"), "https://e.cz/kontakt")
        self.assertEqual(collapse_query("https://e.cz/img/logo.png?v=3"),
                         "https://e.cz/img/logo.png")
        self.assertEqual(collapse_query("https://e.cz/?s=hledat"), "https://e.cz/")

    def test_query_kept_on_script_paths_and_without_query(self):
        self.assertEqual(collapse_query("https://e.cz/index.php?id=5"),
                         "https://e.cz/index.php?id=5")
        self.assertEqual(collapse_query("https://e.cz/detail.aspx?id=5"),
                         "https://e.cz/detail.aspx?id=5")
        self.assertEqual(collapse_query("https://e.cz/o-nas/"), "https://e.cz/o-nas/")

    def test_is_network_error(self):
        dns = {"status": 0, "error": "HTTPSConnectionPool(host='e.cz', port=443): Max retries "
                                     "exceeded with url: /x (Caused by NameResolutionError(...))"}
        self.assertTrue(is_network_error(dns))
        self.assertFalse(is_network_error({"status": 0, "error": "Read timed out. (read timeout=10)"}))
        self.assertFalse(is_network_error({"status": 404, "error": ""}))


# ── check_resources ──────────────────────────────────────────────────────────

class TestCheckResources(unittest.TestCase):
    """probe_url je mockované tabulkou url → (status, size)."""

    def _run(self, results, table, **kw):
        calls = []

        def _fake_probe(session, url, timeout=None):
            calls.append(url)
            status, size = table.get(url, (200, None))
            return {"status": status, "size": size,
                    "error": "" if status else "timeout"}

        with patch.object(links_check, "probe_url", side_effect=_fake_probe), \
             patch.object(links_check, "LINK_CHECK_DELAY", 0):
            report = check_resources(results, BASE, session=MagicMock(), **kw)
        return report, calls

    def test_audited_pages_not_probed_again(self):
        results = [_page(BASE, refs={"links": ["https://example.cz/o-nas/"]}),
                   _page("https://example.cz/o-nas", refs={"links": [BASE]})]
        report, calls = self._run(results, {})
        self.assertEqual(calls, [])
        self.assertEqual(report["known_ok"], 2)
        self.assertEqual(report["broken_links"], [])

    def test_broken_internal_link_reported_with_sources_and_issue(self):
        results = [_page(BASE, refs={"links": ["https://example.cz/mrtva"]}),
                   _page("https://example.cz/b", refs={"links": ["https://example.cz/mrtva/",
                                                                 "https://example.cz/ok"]})]
        report, calls = self._run(results, {"https://example.cz/mrtva": (404, None)})
        self.assertEqual(sorted(calls), ["https://example.cz/mrtva", "https://example.cz/ok"])
        self.assertEqual(len(report["broken_links"]), 1)
        b = report["broken_links"][0]
        self.assertEqual(b["status"], 404)
        self.assertEqual(b["sources"], [BASE, "https://example.cz/b"])
        for r in results:
            issue = next(i for i in r["structure_issues"] if i.type == IssueType.BROKEN_LINK)
            self.assertEqual(issue.items, ["https://example.cz/mrtva (HTTP 404)"])
        self.assertEqual(report["checked_links"], 2)

    def test_external_skipped_unless_flag(self):
        results = [_page(BASE, refs={"links": ["https://jinde.cz/x"],
                                     "images": ["https://cdn.jinde.cz/i.png"]})]
        report, calls = self._run(results, {"https://jinde.cz/x": (404, None)})
        self.assertEqual(calls, [])
        self.assertEqual(report["skipped_external"], 2)
        self.assertEqual(report["broken_links"], [])

        report, calls = self._run(results, {"https://jinde.cz/x": (404, None)},
                                  check_external=True)
        self.assertEqual(sorted(calls), ["https://cdn.jinde.cz/i.png", "https://jinde.cz/x"])
        self.assertEqual(report["skipped_external"], 0)
        self.assertTrue(report["broken_links"][0]["external"])

    def test_external_403_is_not_broken(self):
        results = [_page(BASE, refs={"links": ["https://jinde.cz/x"]})]
        report, _ = self._run(results, {"https://jinde.cz/x": (403, None)}, check_external=True)
        self.assertEqual(report["broken_links"], [])

    def test_images_broken_and_large(self):
        results = [_page(BASE, refs={"images": ["https://example.cz/404.png",
                                                "https://example.cz/big.jpg",
                                                "https://example.cz/ok.jpg"]})]
        report, _ = self._run(results, {"https://example.cz/404.png": (404, None),
                                        "https://example.cz/big.jpg": (200, 700 * 1024),
                                        "https://example.cz/ok.jpg": (200, 10 * 1024)})
        problems = {im["url"]: im for im in report["images"]}
        self.assertEqual(problems["https://example.cz/404.png"]["problem"], "broken")
        self.assertEqual(problems["https://example.cz/big.jpg"]["problem"], "large")
        self.assertEqual(problems["https://example.cz/big.jpg"]["size_kb"], 700)
        self.assertNotIn("https://example.cz/ok.jpg", problems)
        types = {i.type for i in results[0]["structure_issues"]}
        self.assertIn(IssueType.IMG_BROKEN, types)
        self.assertIn(IssueType.IMG_TOO_LARGE, types)
        large = next(i for i in results[0]["structure_issues"] if i.type == IssueType.IMG_TOO_LARGE)
        self.assertEqual(large.items, ["https://example.cz/big.jpg (700 kB)"])
        self.assertEqual(report["checked_images"], 3)

    def test_network_error_counts_as_broken(self):
        results = [_page(BASE, refs={"links": ["https://example.cz/timeout"]})]
        report, _ = self._run(results, {"https://example.cz/timeout": (0, None)})
        self.assertEqual(report["broken_links"][0]["status"], 0)
        issue = results[0]["structure_issues"][0]
        self.assertIn("nedostupné (timeout)", issue.items[0])

    def test_query_variants_collapsed_to_base_page(self):
        """E-shop filtry: 3 varianty kategorie = 1 ověření základní stránky;
        auditovaná kategorie se neověřuje vůbec. index.php?id= zůstává."""
        results = [_page(BASE, refs={"links": ["https://example.cz/kremy/?p13[0]=1",
                                               "https://example.cz/kremy/?p13[0]=2",
                                               "https://example.cz/kremy/?r=cs",
                                               "https://example.cz/akce?page=2",
                                               "https://example.cz/index.php?id=7",
                                               "https://example.cz/logo.png?v=3"]}),
                   _page("https://example.cz/kremy/", refs={})]
        report, calls = self._run(results, {"https://example.cz/akce": (404, None)})
        self.assertEqual(sorted(calls), ["https://example.cz/akce",
                                         "https://example.cz/index.php?id=7",
                                         "https://example.cz/logo.png"])
        self.assertEqual(report["collapsed_query"], 5)
        self.assertEqual(report["known_ok"], 1)
        self.assertEqual([b["url"] for b in report["broken_links"]],
                         ["https://example.cz/akce"])

    def test_max_targets_prefers_most_referenced(self):
        results = [_page(BASE, refs={"links": ["https://example.cz/a", "https://example.cz/b",
                                               "https://example.cz/c"]}),
                   _page("https://example.cz/x", refs={"links": ["https://example.cz/c"]})]
        report, calls = self._run(results, {"https://example.cz/a": (404, None),
                                            "https://example.cz/c": (404, None)},
                                  max_targets=1)
        self.assertEqual(calls, ["https://example.cz/c"])
        self.assertEqual(report["skipped_limit"], 2)
        self.assertEqual(report["checked_links"], 1)
        self.assertEqual([b["url"] for b in report["broken_links"]], ["https://example.cz/c"])

    def test_time_budget_leaves_rest_unverified(self):
        results = [_page(BASE, refs={"links": [f"https://example.cz/p{i}" for i in range(6)]})]
        calls = []

        def _slow_probe(session, url, timeout=None):
            calls.append(url); time.sleep(0.03)
            return {"status": 404, "size": None, "error": ""}

        with patch.object(links_check, "probe_url", side_effect=_slow_probe), \
             patch.object(links_check, "LINK_CHECK_DELAY", 0):
            report = check_resources(results, BASE, session=MagicMock(), workers=1,
                                     max_seconds=0.01)
        self.assertLess(len(calls), 6)
        self.assertEqual(report["skipped_time"], 6 - len(calls))
        self.assertEqual(len(report["broken_links"]), len(calls))
        self.assertEqual(report["aborted"], "")

    def test_dns_error_on_own_domain_is_unverified_not_broken(self):
        results = [_page(BASE, refs={"links": ["https://example.cz/a"],
                                     "images": ["https://example.cz/i.png"]})]
        calls = []

        def _dns_fail(session, url, timeout=None):
            calls.append(url)
            return {"status": 0, "size": None,
                    "error": "Max retries exceeded (Caused by NameResolutionError(...))"}

        with patch.object(links_check, "probe_url", side_effect=_dns_fail), \
             patch.object(links_check, "LINK_CHECK_DELAY", 0):
            report = check_resources(results, BASE, session=MagicMock(), workers=1)
        self.assertEqual(len(calls), 2)
        self.assertEqual(report["broken_links"], [])
        self.assertEqual(report["images"], [])
        self.assertEqual(report["unverified"], 2)
        self.assertEqual(results[0]["structure_issues"], [])

    def test_consecutive_network_errors_abort_phase(self):
        """Výpadek sítě: po LINK_CHECK_ABORT_AFTER chybách za sebou se skončí,
        nic se nehlásí jako chyba a dřívější 404 zůstává."""
        links = ["https://example.cz/real404"] + [f"https://example.cz/p{i}" for i in range(40)]
        results = [_page(BASE, refs={"links": links})]
        calls = []

        def _probe(session, url, timeout=None):
            calls.append(url); time.sleep(0.01)
            if url.endswith("real404"):
                return {"status": 404, "size": None, "error": ""}
            return {"status": 0, "size": None, "error": "Read timed out. (read timeout=10)"}

        with patch.object(links_check, "probe_url", side_effect=_probe), \
             patch.object(links_check, "LINK_CHECK_DELAY", 0), \
             patch.object(links_check, "LINK_CHECK_ABORT_AFTER", 5):
            report = check_resources(results, BASE, session=MagicMock(), workers=1)
        self.assertIn("5 síťových chyb za sebou", report["aborted"])
        self.assertLess(len(calls), 41)
        self.assertEqual([b["url"] for b in report["broken_links"]],
                         ["https://example.cz/real404"])
        self.assertEqual(report["unverified"], 40)
        self.assertEqual(report["skipped_time"], 0)
        issue = results[0]["structure_issues"][0]
        self.assertEqual(issue.count, 1)

    def test_results_without_refs_tolerated(self):
        results = [{"url": BASE, "w3c_category": "ok", "structure_issues": []}]
        report, calls = self._run(results, {})
        self.assertEqual(calls, [])
        self.assertEqual(report["broken_links"], [])


# ── JSON export ──────────────────────────────────────────────────────────────

class TestJsonExport(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_build_json_shape(self):
        results = [_page(BASE, issues=[Issue(IssueType.MISSING_H1)],
                         errors=[{"message": "Bad.", "line": 3}, "plain"]),
                   _page("https://example.cz/x", category="validator_error")]
        data = build_json(results, BASE, source_label="sitemap",
                          link_report={"broken_links": [{"url": "u"}], "checked_links": 5},
                          domain_info={"robots_skipped": True})
        self.assertEqual(data["version"], 1)
        self.assertEqual(data["url"], BASE)
        self.assertEqual(data["summary"]["total"], 2)
        self.assertEqual(data["summary"]["w3c_failed"], 1)
        p0 = data["pages"][0]
        self.assertEqual(p0["issues"][0]["type"], "missing_h1")
        self.assertEqual(p0["issues"][0]["label"], "Chybí <h1> tag")
        self.assertEqual(p0["w3c_errors"], [{"message": "Bad.", "line": 3},
                                            {"message": "plain", "line": None}])
        self.assertEqual(p0["score"], 100 - 15 - 2 * 2)
        self.assertEqual(data["pages"][1]["score"], 0)
        self.assertEqual(data["broken_links"], [{"url": "u"}])
        self.assertEqual(data["links_summary"]["checked_links"], 5)
        self.assertTrue(data["domain"]["robots_skipped"])
        self.assertIsNone(data["comparison"])
        json.dumps(data)   # serializovatelné

    def test_write_and_load_roundtrip_utf8(self):
        data = build_json([_page(BASE, title="Příliš žluťoučký")], BASE)
        path = write_json(data, self.dir / "a.json")
        raw = path.read_text(encoding="utf-8")
        self.assertIn("Příliš žluťoučký", raw)            # ensure_ascii=False
        self.assertEqual(load_previous(path)["pages"][0]["title"], "Příliš žluťoučký")

    def test_write_json_locked_file_saves_with_timestamp(self):
        target = self.dir / "a.json"
        real_write = Path.write_text

        def _write(self_path, *a, **kw):
            if self_path == target:
                raise PermissionError("locked")
            return real_write(self_path, *a, **kw)

        with patch.object(Path, "write_text", _write):
            saved = write_json({"x": 1}, target)
        self.assertNotEqual(saved, target)
        self.assertTrue(saved.name.startswith("a_") and saved.suffix == ".json")

    def test_build_json_path(self):
        xlsx = Path("out") / "ex_validator_20260101_120000.xlsx"
        self.assertEqual(build_json_path(xlsx), Path("out") / "ex_validator_20260101_120000.json")
        self.assertEqual(build_json_path(xlsx, str(self.dir / "v.json")), self.dir / "v.json")
        self.assertEqual(build_json_path(xlsx, str(self.dir)),
                         self.dir / "ex_validator_20260101_120000.json")

    def test_find_previous_prefers_newest_including_keep_stamps(self):
        plain = self.dir / "ex_validator.json"
        old   = self.dir / "ex_validator_20260101_120000.json"
        other = self.dir / "jiny_validator.json"
        for p in (plain, old, other):
            p.write_text("{}", encoding="utf-8")
        now = time.time()
        os.utime(plain, (now - 100, now - 100))
        os.utime(old,   (now, now))
        os.utime(other, (now + 100, now + 100))
        self.assertEqual(find_previous_json(plain), old)
        # nový soubor s časovou značkou (--keep) ještě neexistuje → hledá se vedle
        new_stamped = self.dir / "ex_validator_20260202_120000.json"
        self.assertEqual(find_previous_json(new_stamped), old)
        self.assertIsNone(find_previous_json(self.dir / "nic_validator.json"))

    def test_load_previous_rejects_garbage(self):
        bad = self.dir / "bad.json"
        bad.write_text("{not json", encoding="utf-8")
        self.assertIsNone(load_previous(bad))
        bad.write_text('{"score": 5}', encoding="utf-8")   # chybí pages
        self.assertIsNone(load_previous(bad))
        self.assertIsNone(load_previous(None))


# ── Porovnání ────────────────────────────────────────────────────────────────

class TestCompareRuns(unittest.TestCase):
    def _json(self, pages, score):
        return {"score": score, "generated_at": "2026-09-15T20:55:10",
                "pages": [{"url": u, "issues": [{"label": l} for l in labels],
                           "w3c_errors": [{"message": m, "line": i} for i, m in enumerate(w3c)]}
                          for u, labels, w3c in pages]}

    def test_fixed_new_and_score_delta(self):
        prev = self._json([("https://example.cz/", ["A", "B"], ["E1"]),
                           ("https://example.cz/x", ["C"], [])], 72)
        cur  = self._json([("https://www.example.cz", ["B", "D"], []),      # www + bez lomítka
                           ("https://example.cz/x/", ["C"], ["E2"])], 85)
        c = compare_runs(prev, cur)
        self.assertEqual((c["previous_score"], c["score"], c["delta"]), (72, 85, 13))
        self.assertEqual(c["fixed"], [{"url": "https://www.example.cz", "label": "A"},
                                      {"url": "https://www.example.cz", "label": "W3C: E1"}])
        self.assertEqual(c["new"], [{"url": "https://www.example.cz", "label": "D"},
                                    {"url": "https://example.cz/x/", "label": "W3C: E2"}])
        self.assertEqual((c["fixed_count"], c["new_count"], c["pages_common"]), (2, 2, 2))

    def test_only_common_pages_compared(self):
        prev = self._json([("https://example.cz/stara", ["A"], [])], 50)
        cur  = self._json([("https://example.cz/nova", ["B"], [])], 50)
        c = compare_runs(prev, cur)
        self.assertEqual(c["fixed"], []); self.assertEqual(c["new"], [])
        self.assertEqual((c["pages_added"], c["pages_removed"], c["pages_common"]), (1, 1, 0))

    def test_w3c_line_change_is_not_a_change(self):
        prev = self._json([("https://example.cz/", [], ["E1"])], 90)
        cur  = {"score": 90, "pages": [{"url": "https://example.cz/", "issues": [],
                                        "w3c_errors": [{"message": "E1", "line": 99}]}]}
        c = compare_runs(prev, cur)
        self.assertEqual(c["fixed_count"], 0); self.assertEqual(c["new_count"], 0)

    def test_format_previous_date(self):
        self.assertEqual(format_previous_date("2026-09-15T20:55:10"), "15.09.2026 20:55")
        self.assertEqual(format_previous_date("nesmysl"), "nesmysl")
        self.assertEqual(format_previous_date(""), "?")


# ── Excel – nové sekce ───────────────────────────────────────────────────────

class TestExcelLinksImagesComparison(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.out = Path(self._tmp.name) / "r.xlsx"

    def tearDown(self):
        self._tmp.cleanup()

    def _write(self, **kw):
        from report_excel import write_report
        write_report([_page(BASE)], self.out, BASE, score=80, **kw)
        return _sheet_values(self.out)

    def test_sections_absent_without_link_report(self):
        vals = self._write()
        self.assertFalse(any("NEFUNKČNÍ ODKAZY" in v for v in vals))
        self.assertFalse(any("ZMĚNY OD MINULÉHO" in v for v in vals))

    def test_all_ok_messages_and_external_note(self):
        vals = self._write(link_report={"broken_links": [], "images": [],
                                        "checked_links": 10, "checked_images": 3,
                                        "known_ok": 2, "skipped_external": 7,
                                        "collapsed_query": 120, "skipped_limit": 30,
                                        "aborted": "15 síťových chyb za sebou – výpadek",
                                        "check_external": False})
        self.assertIn("✓ Všechny interní odkazy fungují (12 ověřeno)", vals)
        self.assertIn("✓ Žádné nedostupné ani příliš velké obrázky (3 ověřeno)", vals)
        self.assertTrue(any("7 externích cílů nebylo ověřeno" in v for v in vals))
        self.assertTrue(any("120 URL s parametry" in v for v in vals))
        self.assertTrue(any("30 cílů nad limit" in v for v in vals))
        self.assertTrue(any(v.startswith("⚠ Kontrola odkazů přerušena: 15 síťových") for v in vals))

    def test_broken_links_and_images_rows(self):
        report = {
            "broken_links": [{"url": "http://example.cz/mrtva", "status": 404, "error": "",
                              "external": False, "sources": [BASE, "https://example.cz/b"]},
                             {"url": "https://jinde.cz/x", "status": 0, "error": "timeout",
                              "external": True, "sources": [BASE]}],
            "images": [{"url": "https://example.cz/404.png", "status": 404, "error": "",
                        "size_kb": None, "problem": "broken", "sources": [BASE]},
                       {"url": "https://example.cz/big.jpg", "status": 200, "error": "",
                        "size_kb": 700, "problem": "large", "sources": [BASE]}],
            "checked_links": 2, "checked_images": 2, "known_ok": 0,
            "skipped_external": 0, "check_external": True,
        }
        vals = self._write(link_report=report)
        self.assertIn("https://example.cz/mrtva", vals)              # http → https
        self.assertIn("HTTP 404", vals)
        self.assertIn("https://jinde.cz/x  [externí]", vals)
        self.assertIn("CHYBA", vals)
        self.assertIn(f"{BASE}\nhttps://example.cz/b", vals)
        self.assertIn("NEDOSTUPNÝ", vals); self.assertIn("VELKÝ", vals)
        self.assertIn("700 kB", vals)
        self.assertTrue(any(v.startswith("OBRÁZKY – NEDOSTUPNÉ / VĚTŠÍ NEŽ") for v in vals))

    def test_comparison_rows(self):
        comparison = {"previous_date": "2026-09-15T20:55:10", "previous_score": 72,
                      "score": 80, "delta": 8,
                      "fixed": [{"url": BASE, "label": "Chybí <h1> tag"}],
                      "new": [], "fixed_count": 1, "new_count": 0,
                      "pages_common": 1, "pages_added": 2, "pages_removed": 0}
        vals = self._write(comparison=comparison)
        self.assertIn("Změna od minulého běhu (15.09.2026 20:55)", vals)
        self.assertIn("72 → 80  (+8)", vals)
        self.assertIn("Opraveno od minula (problémů)", vals)
        self.assertIn("Nové problémy od minula", vals)
        self.assertIn("OPRAVENO", vals)
        self.assertIn("Chybí <h1> tag", vals)
        self.assertIn("   – žádné –", vals)
        self.assertTrue(any("nové stránky v auditu: 2" in v for v in vals))


if __name__ == "__main__":
    unittest.main()
